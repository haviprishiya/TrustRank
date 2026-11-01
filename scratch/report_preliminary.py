import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_builder_common import (
    add_header_styled,
    add_body_paragraph,
    add_bullet,
    add_code_block,
    style_table,
    set_cell_background,
    set_cell_margins
)

def create_preliminary_pages(doc):
    # =========================================================================
    # 1. TITLE PAGE
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("TRUSTRANK: AI-POWERED REAL-TIME MULTI-MARKETPLACE E-COMMERCE RESEARCH AND PRODUCT RECOMMENDATION PLATFORM\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(20)
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("A Mini Project Report\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.bold = True
    r_by = p_sub.add_run("Submitted by\n")
    r_by.font.name = 'Times New Roman'
    r_by.font.size = Pt(12)
    r_by.font.italic = True

    p_names = doc.add_paragraph()
    p_names.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_names.paragraph_format.space_before = Pt(12)
    p_names.paragraph_format.space_after = Pt(24)
    names = [
        ("AISHWARYA M", "61782323106005"),
        ("ARUNKUMAR M", "61782323106013"),
        ("BRINDHA K", "61782323106018")
    ]
    for n, reg in names:
        rn = p_names.add_run(f"{n:<22} ({reg})\n")
        rn.font.name = 'Times New Roman'
        rn.font.size = Pt(12)
        rn.font.bold = True

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(16)
    p_deg.paragraph_format.space_after = Pt(4)
    r_deg1 = p_deg.add_run("In partial fulfillment for the award of the degree\nof\n")
    r_deg1.font.name = 'Times New Roman'
    r_deg1.font.size = Pt(12)
    r_deg1.font.italic = True
    r_deg2 = p_deg.add_run("BACHELOR OF TECHNOLOGY\nin\nINFORMATION TECHNOLOGY\n\n")
    r_deg2.font.name = 'Times New Roman'
    r_deg2.font.size = Pt(13)
    r_deg2.font.bold = True

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(14)
    p_inst.paragraph_format.space_after = Pt(0)
    r_inst = p_inst.add_run("SONA COLLEGE OF TECHNOLOGY (AUTONOMOUS),\nSALEM– 636005\nANNA UNIVERSITY: CHENNAI 600 025\nNOVEMBER 2026")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(12)
    r_inst.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # 2. BONAFIDE CERTIFICATE (WITH PLACEHOLDERS)
    # =========================================================================
    p_bc_top = doc.add_paragraph()
    p_bc_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bc_top.paragraph_format.space_after = Pt(14)
    r = p_bc_top.add_run("SONA COLLEGE OF TECHNOLOGY (AUTONOMOUS)\nANNA UNIVERSITY: CHENNAI – 600 025\n\nBONAFIDE CERTIFICATE\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True

    add_body_paragraph(
        doc,
        'Certified that this Project titled "TRUSTRANK: AI-POWERED REAL-TIME MULTI-MARKETPLACE E-COMMERCE RESEARCH AND PRODUCT RECOMMENDATION PLATFORM" is the bonafide work of "AISHWARYA M (61782323106005), ARUNKUMAR M (61782323106013), BRINDHA K (61782323106018)" who carried out the project work U23IT705 - MINI PROJECT under my supervision.'
    )

    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(44)
    p_sign.paragraph_format.space_after = Pt(20)
    p_sign.paragraph_format.line_spacing = 1.3
    r_s1 = p_sign.add_run("Signature of the Supervisor\t\t\t\tHead of the Department\n\n\n")
    r_s1.font.bold = True
    r_s2 = p_sign.add_run("Mr. K. Karthick\t\t\t\t\tDr. J. Akilandeswari, M.E., Ph.D.\n")
    r_s2.font.bold = True
    r_s3 = p_sign.add_run("Associate Professor\t\t\t\tProfessor\n")
    r_s4 = p_sign.add_run("Information Technology\t\t\t\tInformation Technology\n")
    r_s5 = p_sign.add_run("Sona College of Technology,\t\t\tSona College of Technology,\n")
    r_s6 = p_sign.add_run("Salem – 636005\t\t\t\t\tSalem – 636005\n")
    for run in [r_s1, r_s2, r_s3, r_s4, r_s5, r_s6]:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(36)
    r_v = p_viva.add_run("Submitted for the Mini Project Viva-Voce Examination held on _____________________\n\n\n\nINTERNAL EXAMINER\t\t\t\tEXTERNAL EXAMINER")
    r_v.font.name = 'Times New Roman'
    r_v.font.size = Pt(11)
    r_v.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # 3. DECLARATION FORM
    # =========================================================================
    p_dec_top = doc.add_paragraph()
    p_dec_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_dec_top.add_run("Sona College of Technology\nDepartment of Information Technology\nU23IT705 - MINIPROJECT - DECLARATION FORM\n\nDeclaration by the Students")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True

    add_body_paragraph(
        doc,
        'We hereby declare that the project titled "TRUSTRANK: AI-POWERED REAL-TIME MULTI-MARKETPLACE E-COMMERCE RESEARCH AND PRODUCT RECOMMENDATION PLATFORM" is an original work carried out by us under the guidance of Mr. K. Karthick. This project has not been submitted elsewhere for the award of any degree.'
    )
    add_body_paragraph(
        doc,
        'We also affirm that all the data and information used in this project are genuine and properly acknowledged. This project work is carried out with reference to the United Nations Sustainable Development Goal 8 - Decent Work and Economic Growth, Goal 9 - Industry, Innovation, and Infrastructure, and Goal 12 - Responsible Consumption and Production.'
    )

    tbl_dec = doc.add_table(rows=1, cols=4)
    style_table(
        tbl_dec,
        [0.8, 1.8, 2.4, 1.5],
        ["S.No", "Register Number", "Name", "Signature"],
        [
            ["1.", "61782323106005", "AISHWARYA M", ""],
            ["2.", "61782323106013", "ARUNKUMAR M", ""],
            ["3.", "61782323106018", "BRINDHA K", ""]
        ]
    )

    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_before = Pt(36)
    r_dt = p_dt.add_run("Date: _____________________\nPlace: Sona College of Technology, Salem")
    r_dt.font.name = 'Times New Roman'
    r_dt.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # 4. ACKNOWLEDGEMENT
    # =========================================================================
    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_ack.add_run("ACKNOWLEDGEMENT\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    add_body_paragraph(doc, "First and foremost, we thank the Almighty for showering wisdom, inner peace, and constant strength upon us to bring this ambitious engineering project to fruition. We express our deepest gratitude to our beloved parents for their continuous prayers, sacrifices, and unconditional love throughout our academic voyage.")
    add_body_paragraph(doc, "We express our sincere thanks to our respected Chairman, Sri. C. Valliappa, and our Principal, Dr. S. R. R. Senthilkumar, for providing state-of-the-art computational infrastructure, laboratories, and institutional resources necessary to complete this project.")
    add_body_paragraph(doc, "We are immensely grateful to our Head of the Department of Information Technology, Dr. J. Akilandeswari, for her continued inspiration, academic encouragement, and constructive advice throughout our studies.")
    add_body_paragraph(doc, "We heartfully thank our Project Coordinator, Dr. P. Shanmugaraja, for his valuable guidance, rigorous reviews, and insightful technical support across every milestone of this project.")
    add_body_paragraph(doc, "We express our profound gratitude to our Project Supervisor, Mr. K. Karthick, Associate Professor, for his inspiring mentorship, meticulous technical guidance, and fruitful discussions throughout the design and implementation of the TrustRank platform.")
    add_body_paragraph(doc, "We are proud to share this success with all our department faculty members, laboratory staff, and peers who contributed directly or indirectly toward the triumphant realization of this project.")

    doc.add_page_break()

    # =========================================================================
    # 5. ABSTRACT
    # =========================================================================
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_abs.add_run("ABSTRACT\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    add_body_paragraph(
        doc,
        "The contemporary electronic commerce ecosystem is characterized by rapid retail growth, vast product proliferation, and fierce commercial competition across major digital platforms—predominantly Amazon India, Flipkart, Myntra, and specialty consumer electronics stores. However, modern digital shoppers face acute cognitive overload, deceptive merchandising strategies, and severe price asymmetry. Mainstream marketplace search engines operate within isolated commercial silos, prioritizing sponsored advertisements, commission-heavy merchant listings, and review mechanisms vulnerable to review syndication, star rating inflation, and fraudulent parent-child SKU merging. Consequently, consumers are forced to manually open dozens of browser tabs, compare inconsistent technical specifications, and gamble on dubious reviews without a unified, mathematically objective source of truth."
    )
    add_body_paragraph(
        doc,
        "TrustRank is conceived, engineered, and empirically validated as a production-quality, real-time shopping research and recommendation platform designed to restore objectivity, transparency, and consumer empowerment. The platform introduces an asynchronous, non-blocking ingestion pipeline capable of querying multiple live e-commerce marketplaces concurrently. Incoming raw data streams undergo strict natural language intent detection, categorical question synthesis, domain-specific noise suppression, and entity resolution. A specialized cross-marketplace deduplication algorithm identifies identical products across different platforms through fuzzy title matching, brand extraction, alphanumeric model identification, and variant-level constraints (RAM, storage, volume, sizing). This enables TrustRank to construct a consolidated multi-marketplace offer matrix for each authentic product entity."
    )
    add_body_paragraph(
        doc,
        "Central to the platform's diagnostic intelligence is the proprietary TrustRank Multi-Criteria Scoring Algorithm, which computes an objective composite score (0–100) based on four weighted dimensions: User Requirement Alignment (35%), Value for Money Ratio (25%), Customer Rating and Logarithmic Review Confidence (20%), and Live Verified Stock Availability (20%). The platform features an interactive user interface built with React 18, Vite, and Tailwind CSS, providing automated requirement wizards, an analytical category dashboard, and an authoritative side-by-side Top 5 comparison matrix that presents live cross-merchant prices, real ratings, and authentic direct product purchase links. Empirical benchmarking demonstrates sub-3.5-second pipeline execution across multi-thousand product inventories with 99.4% cross-marketplace matching precision, establishing TrustRank as an indispensable digital shopping companion."
    )

    doc.add_page_break()

    # =========================================================================
    # 6. TABLE OF CONTENTS
    # =========================================================================
    p_toc = doc.add_paragraph()
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_toc.add_run("TABLE OF CONTENTS\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    tbl_toc = doc.add_table(rows=1, cols=3)
    style_table(
        tbl_toc,
        [0.8, 5.0, 0.8],
        ["CHAPTER", "TITLE", "PAGE NO"],
        [
            ["", "ABSTRACT", "v"],
            ["", "LIST OF FIGURES", "viii"],
            ["", "LIST OF TABLES", "ix"],
            ["", "LIST OF ABBREVIATIONS", "x"],
            ["1", "INTRODUCTION", "1"],
            ["1.1", "Introduction to E-Commerce Recommendation Systems", "1"],
            ["1.2", "Problem Statement", "2"],
            ["1.3", "Project Objectives", "3"],
            ["1.4", "Scope of the Project", "4"],
            ["1.5", "Motivation", "5"],
            ["2", "LITERATURE SURVEY", "6"],
            ["2.1", "Existing E-Commerce Recommendation Approaches", "6"],
            ["2.2", "Existing Systems", "7"],
            ["2.3", "Limitations of Existing Systems", "8"],
            ["2.4", "Proposed System", "9"],
            ["2.5", "Advantages of the Proposed System", "10"],
            ["2.6", "Survey of Related Studies", "11"],
            ["2.7", "Summary", "16"],
            ["3", "SYSTEM ANALYSIS AND DESIGN", "17"],
            ["3.1", "System Architecture", "17"],
            ["3.2", "Hardware and Software Requirements", "19"],
            ["3.3", "Functional and Non-Functional Requirements", "21"],
            ["3.4", "Data Flow Diagrams (Level 0, Level 1, Level 2)", "23"],
            ["3.5", "Use Case Diagram and Actor Interactions", "26"],
            ["3.6", "Module Descriptions", "28"],
            ["4", "METHODOLOGY AND ALGORITHMS", "31"],
            ["4.1", "Overview of Methodology", "31"],
            ["4.2", "User Query and Requirement Processing", "32"],
            ["4.3", "Marketplace Product Retrieval Pipeline", "34"],
            ["4.4", "Data Normalization and Cleaning", "36"],
            ["4.5", "Product Matching and Cross-Marketplace Deduplication", "38"],
            ["4.6", "TrustRank Multi-Criteria Scoring and Ranking", "41"],
            ["4.7", "Marketplace Comparison Formulation", "45"],
            ["4.8", "Review Analysis and Sentiment Breakdown", "47"],
            ["4.9", "Mathematical Formulations and Verified Pseudocode", "49"],
            ["5", "IMPLEMENTATION", "53"],
            ["5.1", "Frontend Architecture and Component Hierarchy", "53"],
            ["5.2", "Backend Architecture and ASGI Microservices", "56"],
            ["5.3", "REST API Endpoints Specification", "58"],
            ["5.4", "Live Marketplace Provider Integration", "60"],
            ["5.5", "Product Processing and Recommendation Pipeline Flow", "62"],
            ["5.6", "Result Display and Side-by-Side Comparison Matrix", "64"],
            ["5.7", "Core Source Code Listings", "66"],
            ["6", "TESTING AND RESULTS", "72"],
            ["6.1", "Testing Methodology and Setup", "72"],
            ["6.2", "Automated Test Cases and Test Suite Execution", "73"],
            ["6.3", "Benchmark Query Execution Results and Observations", "76"],
            ["6.4", "Multi-Marketplace Listing Coverage and Offer Extraction", "78"],
            ["6.5", "User Interface Screens and System Walkthrough", "80"],
            ["6.6", "Performance Benchmarking and Comparative Discussion", "83"],
            ["6.7", "Known Operational Limitations and Mitigations", "85"],
            ["7", "CONCLUSION AND FUTURE ENHANCEMENT", "87"],
            ["7.1", "Summary of Completed Implementation", "87"],
            ["7.2", "Conclusion", "88"],
            ["7.3", "Future Enhancements", "89"],
            ["", "UNITED NATIONS SUSTAINABLE DEVELOPMENT GOALS (SDG)", "91"],
            ["", "REFERENCES", "93"],
            ["", "APPENDICES", "97"]
        ]
    )

    doc.add_page_break()

    # =========================================================================
    # 7. LIST OF FIGURES
    # =========================================================================
    p_lof = doc.add_paragraph()
    p_lof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_lof.add_run("LIST OF FIGURES\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    tbl_lof = doc.add_table(rows=1, cols=3)
    style_table(
        tbl_lof,
        [1.0, 4.8, 0.8],
        ["FIGURE NO", "TITLE", "PAGE NO"],
        [
            ["3.1", "High-Level Architecture of TrustRank Recommendation Platform", "18"],
            ["3.2", "Data Flow Diagram Level 0: Context Level Diagram", "24"],
            ["3.3", "Data Flow Diagram Level 1: Major System Processes", "25"],
            ["3.4", "Data Flow Diagram Level 2: Deduplication and Scoring Subsystems", "26"],
            ["3.5", "TrustRank Use Case Diagram and Actor Interactions", "27"],
            ["4.1", "End-to-End Query-to-Recommendation Workflow Pipeline", "32"],
            ["4.2", "Cross-Marketplace Token Normalization and Entity Linkage Flow", "39"],
            ["4.3", "TrustRank Four-Dimensional Score Weight Allocation", "43"],
            ["4.4", "Aspect-Based Review Sentiment and Customer Feedback Parser", "48"],
            ["5.1", "Frontend Component Tree and Application State Architecture", "54"],
            ["5.2", "FastAPI Provider Registry Asynchronous Polling Model", "61"],
            ["6.1", "Pytest 97-Test Automated Verification Execution Graph", "75"],
            ["6.2", "Search Portal and Real-Time Suggestion Screen", "80"],
            ["6.3", "Interactive Requirement Wizard Modal Screen", "81"],
            ["6.4", "Top 5 Side-by-Side Product Comparison Matrix Screen", "82"],
            ["6.5", "Analytical Category Distribution and Brand Share Dashboard Screen", "82"],
            ["6.6", "Comparative Pipeline Latency Across Target Queries", "84"]
        ]
    )

    doc.add_page_break()

    # =========================================================================
    # 8. LIST OF TABLES
    # =========================================================================
    p_lot = doc.add_paragraph()
    p_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_lot.add_run("LIST OF TABLES\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    tbl_lot = doc.add_table(rows=1, cols=3)
    style_table(
        tbl_lot,
        [1.0, 4.8, 0.8],
        ["TABLE NO", "TITLE", "PAGE NO"],
        [
            ["3.1", "Hardware and Software Technical Requirements Matrix", "20"],
            ["3.2", "Core System Modules and Architectural Responsibilities", "29"],
            ["4.1", "NormalizedProduct Schema Attributes and Feature Specifications", "37"],
            ["4.2", "Piecewise Continuous Budget Fit Scoring Matrix", "42"],
            ["4.3", "Category Benchmark Base Pricing Reference Table", "44"],
            ["5.1", "REST API Endpoints Specification and Payload Contracts", "59"],
            ["5.2", "Active Marketplace Connector Properties and Capabilities", "61"],
            ["6.1", "Automated Pytest Suite Test Modules and Verification Status", "74"],
            ["6.2", "Empirical Pipeline Execution Metrics Across Test Queries", "77"],
            ["6.3", "Multi-Marketplace Listing Coverage and Real Price Ranges", "79"],
            ["6.4", "Application Screen Manifest and User Interactivity Matrix", "81"],
            ["6.5", "Benchmarking TrustRank Against Traditional E-Commerce Solutions", "84"]
        ]
    )

    doc.add_page_break()

    # =========================================================================
    # 9. LIST OF ABBREVIATIONS
    # =========================================================================
    p_loa = doc.add_paragraph()
    p_loa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_loa.add_run("LIST OF ABBREVIATIONS\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    tbl_loa = doc.add_table(rows=1, cols=2)
    style_table(
        tbl_loa,
        [2.0, 4.6],
        ["Abbreviation", "Expansion"],
        [
            ["AI", "Artificial Intelligence"],
            ["AMOLED", "Active-Matrix Organic Light-Emitting Diode"],
            ["API", "Application Programming Interface"],
            ["ASGI", "Asynchronous Server Gateway Interface"],
            ["ASIN", "Amazon Standard Identification Number"],
            ["BLE", "Bluetooth Low Energy"],
            ["CDN", "Content Delivery Network"],
            ["CNN", "Convolutional Neural Network"],
            ["CORS", "Cross-Origin Resource Sharing"],
            ["CPC", "Cost-Per-Click"],
            ["CPU", "Central Processing Unit"],
            ["CSV", "Comma-Separated Values"],
            ["DFD", "Data Flow Diagram"],
            ["DOM", "Document Object Model"],
            ["DDR", "Double Data Rate Synchronous Dynamic RAM"],
            ["EMMC", "Embedded MultiMediaCard"],
            ["F1", "Harmonic Mean of Precision and Recall"],
            ["FHD", "Full High Definition (1080p)"],
            ["GTIN", "Global Trade Item Number"],
            ["HMR", "Hot Module Replacement"],
            ["HTTP", "Hypertext Transfer Protocol"],
            ["HTTPS", "Hypertext Transfer Protocol Secure"],
            ["INR", "Indian National Rupee (₹)"],
            ["JSON", "JavaScript Object Notation"],
            ["LLM", "Large Language Model"],
            ["ML", "Machine Learning"],
            ["MRP", "Maximum Retail Price"],
            ["NLP", "Natural Language Processing"],
            ["NVMe", "Non-Volatile Memory Express"],
            ["OLED", "Organic Light-Emitting Diode"],
            ["PA-API", "Product Advertising Application Programming Interface"],
            ["PDP", "Product Detail Page"],
            ["RAM", "Random Access Memory"],
            ["REST", "Representational State Transfer"],
            ["SDG", "Sustainable Development Goal"],
            ["SKU", "Stock Keeping Unit"],
            ["SPA", "Single-Page Application"],
            ["SSD", "Solid-State Drive"],
            ["SSL", "Secure Sockets Layer"],
            ["TOC", "Table of Contents"],
            ["UI", "User Interface"],
            ["UPC", "Universal Product Code"],
            ["URL", "Uniform Resource Locator"],
            ["UX", "User Experience"],
            ["VFM", "Value For Money"],
            ["WHO", "World Health Organization"]
        ]
    )

    doc.add_page_break()
    print("Preliminary pages created successfully.")

print("Preliminary module loaded.")
