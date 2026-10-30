import os
import sys
import time
from docx import Document
from docx.shared import Inches, Pt, RGBColor

# Ensure scratch directories are in python search path
current_dir = os.path.dirname(os.path.abspath(__file__))
if current_dir not in sys.path:
    sys.path.insert(0, current_dir)

brain_scratch = r"C:\Users\haviprishiya\.gemini\antigravity\brain\8bfbb5eb-0e40-4c21-930d-57cee7d60533\scratch"
if os.path.exists(brain_scratch) and brain_scratch not in sys.path:
    sys.path.insert(0, brain_scratch)

from report_builder_common import (
    add_header_styled,
    add_body_paragraph,
    add_bullet,
    add_code_block,
    style_table,
    set_cell_background,
    set_cell_margins
)
from report_preliminary import create_preliminary_pages
from report_chapters_1_2_3 import build_chapter_1, build_chapter_2, build_chapter_3
from report_chapters_4_5 import build_chapter_4, build_chapter_5
from report_chapters_6_7_sdg_ref import (
    build_chapter_6,
    build_chapter_7,
    build_sdg_and_references,
    build_appendices
)

def build_full_report(output_paths):
    print("=" * 80)
    print("BUILDING TRUSTRANK FINAL ACADEMIC PROJECT REPORT")
    print("=" * 80)
    start_total = time.time()
    
    doc = Document()
    
    # Configure standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

    print("Step 1: Assembling Preliminary Pages (Title, Certificate, Declaration, Acknowledgement, Abstract, TOC, LOF, LOT, Abbreviations)...")
    create_preliminary_pages(doc)

    print("Step 2: Assembling Chapter 1 – Introduction...")
    build_chapter_1(doc)

    print("Step 3: Assembling Chapter 2 – Literature Survey...")
    build_chapter_2(doc)

    print("Step 4: Assembling Chapter 3 – System Analysis and Design...")
    build_chapter_3(doc)

    print("Step 5: Assembling Chapter 4 – Methodology and Algorithms...")
    build_chapter_4(doc)

    print("Step 6: Assembling Chapter 5 – Implementation...")
    build_chapter_5(doc)

    print("Step 7: Assembling Chapter 6 – Testing and Results...")
    build_chapter_6(doc)

    print("Step 8: Assembling Chapter 7 – Conclusion and Future Enhancement...")
    build_chapter_7(doc)

    print("Step 9: Assembling United Nations SDG and References...")
    build_sdg_and_references(doc)

    print("Step 10: Assembling Appendices and Source Code Listings...")
    build_appendices(doc)

    # Save to all requested file paths
    print("\nSaving document...")
    for pth in output_paths:
        target_dir = os.path.dirname(os.path.abspath(pth))
        os.makedirs(target_dir, exist_ok=True)
        doc.save(pth)
        file_size_kb = os.path.getsize(pth) / 1024
        print(f" -> Successfully saved: {pth} ({file_size_kb:.2f} KB)")

    print(f"\nAssembly completed successfully in {time.time() - start_total:.2f} seconds!")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print("=" * 80)

if __name__ == "__main__":
    targets = [
        r"c:\Users\haviprishiya\Downloads\update new\update\update\TrustRank_Final_Project_Report.docx",
        r"c:\Users\haviprishiya\Downloads\update new\TrustRank_Final_Project_Report.docx",
        r"C:\Users\haviprishiya\.gemini\antigravity\brain\8bfbb5eb-0e40-4c21-930d-57cee7d60533\TrustRank_Final_Project_Report.docx"
    ]
    build_full_report(targets)
