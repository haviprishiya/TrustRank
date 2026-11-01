import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6, italic_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.font.name = 'Times New Roman'
        r_it.font.size = Pt(12)
        r_it.font.italic = True
        r_it.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, title, desc):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_t = p.add_run(title + ": ")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)
    r_d = p.add_run(desc)
    r_d.font.name = 'Times New Roman'
    r_d.font.size = Pt(12)
    r_d.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(11)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(30, 58, 138)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

def style_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    for row_idx, data in enumerate(rows_data):
        row_cells = table.add_row().cells
        bg_color = "F1F5F9" if row_idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(data):
            row_cells[i].text = str(val)
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=90, bottom=90, left=120, right=120)
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)

    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)
